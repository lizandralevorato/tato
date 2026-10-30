"""Linhas e arquivos (CSV e Excel) das planilhas do painel."""

from datetime import date, datetime
from io import BytesIO

from django.utils import timezone
from openpyxl import Workbook
from openpyxl.styles import Font
from openpyxl.utils import get_column_letter

from apps.agenda.formatacao import hora

from . import relatorio_colaboradores as relatorio

FORMATO_DATA = "dd/mm/yyyy"
FORMATO_DATA_HORA = "dd/mm/yyyy hh:mm"


def _local(momento):
    return timezone.localtime(momento).replace(tzinfo=None) if momento else None


def _sim_nao(valor):
    return "Sim" if valor else "Não"


def _vazio(valor):
    return "" if valor is None else valor


def texto(valor):
    """Valor de uma linha no formato do CSV."""
    if isinstance(valor, datetime):
        return valor.strftime("%d/%m/%Y %H:%M")
    if isinstance(valor, date):
        return valor.strftime("%d/%m/%Y")
    return valor


CABECALHO_AGENDAMENTOS = [
    "Data", "Hora", "Massoterapeuta", "Colaborador", "E-mail", "Situação", "Origem", "Agendado em",
    "Cancelado por", "Fez check-in", "Check-in em", "Chegou em", "Cancelado em", "Encaixe em vaga liberada",
]
FORMATOS_AGENDAMENTOS = {1: FORMATO_DATA, 8: FORMATO_DATA_HORA, 11: FORMATO_DATA_HORA, 12: FORMATO_DATA_HORA, 13: FORMATO_DATA_HORA}


def linhas_agendamentos(agendamentos):
    """`agendamentos` vem de `relatorio_colaboradores.reservas` (traz `em_vaga_liberada`)."""
    return [
        [
            _local(ag.inicio).date(), hora(ag.inicio), ag.massoterapeuta.nome, ag.colaborador.nome,
            ag.colaborador.email, ag.get_status_display(), ag.get_origem_display(), _local(ag.criado_em),
            ag.cancelado_por.nome if ag.cancelado_por else "", _sim_nao(ag.confirmado_em), _local(ag.confirmado_em),
            _local(ag.chegada_em), _local(ag.cancelado_em),
            _sim_nao(ag.origem == ag.Origem.ENCAIXE and ag.em_vaga_liberada),
        ]
        for ag in agendamentos
    ]


CABECALHO_RESUMO = [
    "Massoterapeuta", "Horários oferecidos", "Reservados", "Ocupação (%)", "Atendidos", "Faltas",
    "Faltas (%)", "Cancelados", "Liberados sem check-in", "Pelo administrador ou encaixe",
]


def linhas_resumo(dados):
    if dados is None:
        return []
    return [
        [
            linha.rotulo, linha.oferecidos, linha.reservados, _vazio(linha.ocupacao), linha.realizados,
            linha.faltas, _vazio(linha.taxa_faltas), linha.cancelados, linha.liberados, linha.encaixes,
        ]
        for linha in [*dados["por_massoterapeuta"], dados["geral"]]
    ]


CABECALHO_POR_COLABORADOR = [
    "Colaborador", "E-mail", "Reservas", "Check-ins feitos", "Check-in (%)", "Atendidos", "Faltas",
    "Faltas (%)", "Cancelados", "Liberados sem check-in", "Encaixes feitos", "Encaixes em vaga liberada",
]


def _linhas_por_colaborador(atividades):
    return [
        [
            a.colaborador.nome, a.colaborador.email, a.reservados, a.checkins, _vazio(a.taxa_checkin),
            a.realizados, a.faltas, _vazio(a.taxa_faltas), a.cancelados, a.liberados,
            a.encaixes, a.encaixes_em_vaga_liberada,
        ]
        for a in atividades
    ]


CABECALHO_POR_MES = ["Colaborador", "Mês", "Reservas", "Atendidos", "Faltas", "Cancelados", "Liberados sem check-in"]


def _linhas_por_mes(atividades):
    linhas = []
    for a in atividades:
        for (ano, mes), contagem in sorted(a.por_mes.items()):
            cancelados = contagem["cancelado"]
            liberados = contagem["expirado"]
            linhas.append(
                [
                    a.colaborador.nome, f"{mes:02d}/{ano}", sum(contagem.values()) - cancelados - liberados,
                    contagem["realizado"], contagem["falta"], cancelados, liberados,
                ]
            )
    return linhas


def _aba(planilha, titulo, cabecalho, linhas, formatos=None):
    aba = planilha.create_sheet(titulo)
    aba.append(cabecalho)
    for celula in aba[1]:
        celula.font = Font(bold=True)
    for linha in linhas:
        aba.append(linha)
    for linha in aba.iter_rows(min_row=2):
        for celula in linha:
            # Nomes vindos do cadastro nunca podem virar fórmula.
            if isinstance(celula.value, str):
                celula.data_type = "s"
            formato = (formatos or {}).get(celula.column)
            if formato:
                celula.number_format = formato
    for indice, titulo_coluna in enumerate(cabecalho, start=1):
        tamanhos = [len(str(titulo_coluna))] + [len(str(l[indice - 1])) for l in linhas if l[indice - 1] is not None]
        aba.column_dimensions[get_column_letter(indice)].width = min(max(tamanhos) + 2, 50)
    aba.freeze_panes = "A2"
    aba.auto_filter.ref = aba.dimensions


def _bytes(planilha):
    arquivo = BytesIO()
    planilha.save(arquivo)
    return arquivo.getvalue()


def xlsx_agendamentos(agendamentos):
    """Reservas do período, mais os totais de cada colaborador (geral e por mês)."""
    agendamentos = list(agendamentos)
    atividades = relatorio.atividade_por_colaborador(agendamentos)
    planilha = Workbook()
    planilha.remove(planilha.active)
    _aba(planilha, "Reservas", CABECALHO_AGENDAMENTOS, linhas_agendamentos(agendamentos), FORMATOS_AGENDAMENTOS)
    _aba(planilha, "Por colaborador", CABECALHO_POR_COLABORADOR, _linhas_por_colaborador(atividades))
    _aba(planilha, "Por mês", CABECALHO_POR_MES, _linhas_por_mes(atividades))
    return _bytes(planilha)


def xlsx_resumo(dados):
    planilha = Workbook()
    planilha.remove(planilha.active)
    _aba(planilha, "Por massoterapeuta", CABECALHO_RESUMO, linhas_resumo(dados))
    return _bytes(planilha)
